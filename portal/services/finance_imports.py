"""Normalized, non-destructive bank statement ingestion for ArenaLine finance."""
import csv, hashlib, io, json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from openpyxl import load_workbook
from portal.model_modules.finance import BankImportBatch, ImportedBankTransaction

REQUIRED=("date","description")
def _text(value): return "" if value is None else str(value).strip()
def _date(value):
    if isinstance(value,datetime):return value.date()
    if isinstance(value,date):return value
    text=_text(value)
    for fmt in ("%Y-%m-%d","%m/%d/%Y","%m/%d/%y"):
        try:return datetime.strptime(text,fmt).date()
        except ValueError:pass
    raise ValidationError(f"Unsupported transaction date: {text}")
def _money(value):
    text=_text(value).replace("$","").replace(",","").replace("(","-").replace(")","")
    try:return Decimal(text)
    except InvalidOperation as exc:raise ValidationError(f"Invalid amount: {value}") from exc
def _rows_csv(data):
    text=data.decode("utf-8-sig") if isinstance(data,bytes) else data
    return list(csv.DictReader(io.StringIO(text)))
def _rows_xlsx(data):
    book=load_workbook(io.BytesIO(data),read_only=True,data_only=True);sheet=book.active
    rows=sheet.iter_rows(values_only=True);headers=[_text(v) for v in next(rows,())]
    return [dict(zip(headers,row)) for row in rows]
def parse_rows(data,file_type):
    if file_type=="csv":return _rows_csv(data)
    if file_type=="xlsx":return _rows_xlsx(data)
    raise ValidationError("Only CSV and XLSX bank imports are supported.")
def normalize_row(raw,mapping):
    for key in REQUIRED:
        if not mapping.get(key):raise ValidationError(f"Import mapping requires '{key}'.")
    tx_date=_date(raw.get(mapping["date"]));description=_text(raw.get(mapping["description"]))
    if not description:raise ValidationError("Transaction description is required.")
    if mapping.get("amount"):
        signed=_money(raw.get(mapping["amount"]))
        direction=ImportedBankTransaction.Direction.CREDIT if signed>=0 else ImportedBankTransaction.Direction.DEBIT;amount=abs(signed)
    else:
        credit=_money(raw.get(mapping.get("credit"))) if mapping.get("credit") and _text(raw.get(mapping["credit"])) else Decimal("0")
        debit=_money(raw.get(mapping.get("debit"))) if mapping.get("debit") and _text(raw.get(mapping["debit"])) else Decimal("0")
        if bool(credit)==bool(debit):raise ValidationError("Each row must contain either a credit or debit amount.")
        direction=ImportedBankTransaction.Direction.CREDIT if credit else ImportedBankTransaction.Direction.DEBIT;amount=abs(credit or debit)
    if amount<=0:raise ValidationError("Imported transaction amount must be greater than zero.")
    normalized={"transaction_date":tx_date,"posted_date":_date(raw.get(mapping["posted_date"])) if mapping.get("posted_date") and _text(raw.get(mapping["posted_date"])) else None,"amount":amount,"direction":direction,"description":description,"reference":_text(raw.get(mapping.get("reference"))) if mapping.get("reference") else "","external_id":_text(raw.get(mapping.get("external_id"))) if mapping.get("external_id") else ""}
    fingerprint_payload="|".join([str(normalized["transaction_date"]),str(amount),direction,description.lower(),normalized["reference"],normalized["external_id"]])
    normalized["row_fingerprint"]=hashlib.sha256(fingerprint_payload.encode()).hexdigest()
    return normalized

@transaction.atomic
def stage_bank_import(*,team,financial_account,source_name,data,file_type,column_mapping,profile=None):
    if financial_account.team_id!=team.id:raise ValidationError("Financial account must belong to the importing organization.")
    source_bytes=data if isinstance(data,bytes) else data.encode()
    source_fingerprint=hashlib.sha256(source_bytes).hexdigest()
    if BankImportBatch.objects.filter(team=team,financial_account=financial_account,source_fingerprint=source_fingerprint).exists():raise ValidationError("This bank statement has already been imported to this financial account.")
    rows=parse_rows(data,file_type)
    batch=BankImportBatch(team=team,financial_account=financial_account,profile=profile,finance_domain=financial_account.finance_domain,source_name=source_name,source_fingerprint=source_fingerprint);batch.full_clean();batch.save()
    seen=set()
    for raw in rows:
        normalized=normalize_row(raw,column_mapping);fingerprint=normalized["row_fingerprint"]
        if fingerprint in seen:continue
        seen.add(fingerprint)
        row=ImportedBankTransaction(batch=batch,raw_data={str(k):_text(v) for k,v in raw.items()},**normalized);row.full_clean();row.save()
    return batch
